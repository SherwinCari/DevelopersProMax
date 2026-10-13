"""Tests for grade photo extraction parsing (grade-photo-import Tasks 1-2).

These exercise the pure parsing/validation path only; no live vision model.
"""
from __future__ import annotations

from vision import GradeExtraction, _parse_extraction, extract_grades


def test_well_formed_json():
    r = _parse_extraction('{"gwa": 88, "math": 90, "english": 85, "shs_track": "STEM"}')
    assert r.ok is True
    assert r.gwa == 88 and r.math == 90 and r.english == 85
    assert r.shs_track == "STEM"
    assert set(r.filled_fields) == {"GWA", "Math", "English", "SHS strand"}


def test_json_embedded_in_prose():
    text = (
        "Sure! Here are the grades I found:\n"
        '{"gwa": 91.5, "math": 87, "english": 84, "shs_track": "ICT"}\n'
        "Let me know if you need anything else."
    )
    r = _parse_extraction(text)
    assert r.ok is True
    assert r.gwa == 91.5 and r.math == 87 and r.english == 84
    assert r.shs_track == "ICT (TVL)"  # mapped to catalog track


def test_out_of_range_values_dropped_with_note():
    r = _parse_extraction('{"gwa": 150, "math": 90, "english": 60, "shs_track": null}')
    assert r.gwa is None          # 150 out of range -> dropped
    assert r.english is None       # 60 below 75 -> dropped
    assert r.math == 90
    assert r.ok is True            # at least one field read
    assert "out-of-range" in r.notes.lower()


def test_nulls_are_handled():
    r = _parse_extraction('{"gwa": null, "math": null, "english": null, "shs_track": null}')
    assert r.ok is False
    assert r.gwa is None and r.math is None and r.english is None
    assert "manually" in r.notes.lower()


def test_garbage_response_is_fallback():
    r = _parse_extraction("I'm not sure, the image is blurry.")
    assert r.ok is False
    assert "manually" in r.notes.lower()


def test_empty_response_is_fallback():
    r = _parse_extraction("")
    assert r.ok is False


def test_unknown_track_becomes_none():
    r = _parse_extraction('{"gwa": 90, "math": 90, "english": 90, "shs_track": "Underwater Basket Weaving"}')
    assert r.shs_track is None
    assert r.ok is True  # grades still read


def test_extract_grades_rejects_empty_bytes():
    r = extract_grades(b"")
    assert r.ok is False
    assert isinstance(r, GradeExtraction)


def test_extract_grades_rejects_oversized_image():
    big = b"x" * (10 * 1024 * 1024 + 1)
    r = extract_grades(big)
    assert r.ok is False
    assert "10 mb" in r.notes.lower()


# ---- document (PDF / DOCX) support ----
import io  # noqa: E402

from vision import (  # noqa: E402
    _docx_to_text,
    _ext_of,
    _pdf_to_text,
    extract_grades_from_upload,
)


def test_ext_detection():
    assert _ext_of("grades.PDF") == "pdf"
    assert _ext_of("report.docx") == "docx"
    assert _ext_of("photo.JPG") == "jpg"
    assert _ext_of("noext") == ""


def test_upload_empty_bytes():
    r = extract_grades_from_upload(b"", "x.pdf")
    assert r.ok is False


def test_upload_unsupported_type():
    r = extract_grades_from_upload(b"something", "notes.txt")
    assert r.ok is False
    assert "unsupported" in r.notes.lower()


def test_upload_oversized_doc():
    big = b"%PDF-" + b"x" * (20 * 1024 * 1024 + 1)
    r = extract_grades_from_upload(big, "big.pdf")
    assert r.ok is False
    assert "20 mb" in r.notes.lower()


def test_docx_text_extraction():
    import docx

    d = docx.Document()
    d.add_paragraph("Senior High School Report Card")
    d.add_paragraph("GWA: 90   Math: 88   English: 86   Strand: STEM")
    buf = io.BytesIO()
    d.save(buf)
    text = _docx_to_text(buf.getvalue())
    assert "GWA" in text and "STEM" in text


def test_pdf_text_extraction_roundtrip():
    from pypdf import PdfWriter

    # A blank PDF has little/no text -> _pdf_to_text returns a short string.
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    text = _pdf_to_text(buf.getvalue())
    assert isinstance(text, str)  # does not raise; scanned-style PDFs yield little text
